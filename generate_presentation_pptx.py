"""
Generate SIH 2026 Presentation PPTX for Problem Statement SIH26117
Sovereign On-Premise Agentic AI Workbench - MRPL
Using python-pptx with a 16:9 widescreen format and custom card/table layouts.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # Blank slide

    # Color Palette Constants
    C_SLATE_900 = RGBColor(15, 23, 42)     # #0F172A
    C_SLATE_800 = RGBColor(30, 41, 59)     # #1E293B
    C_SLATE_700 = RGBColor(51, 65, 85)     # #334155
    C_SLATE_600 = RGBColor(71, 85, 105)    # #475569
    C_SLATE_500 = RGBColor(100, 116, 139)  # #64748B
    C_SLATE_300 = RGBColor(203, 213, 225)  # #CBD5E1
    C_SLATE_100 = RGBColor(241, 245, 249)  # #F1F5F9
    C_SLATE_50  = RGBColor(248, 250, 252)  # #F8FAFC
    
    C_AMBER_600 = RGBColor(217, 119, 6)    # #D97706
    C_AMBER_500 = RGBColor(245, 158, 11)   # #F59E0B
    C_AMBER_100 = RGBColor(254, 243, 199)  # #FEF3C7
    
    C_GREEN_700 = RGBColor(21, 128, 61)    # #15803D
    C_GREEN_600 = RGBColor(22, 163, 74)    # #16A34A
    C_GREEN_50  = RGBColor(240, 253, 244)  # #F0FDF4
    C_GREEN_200 = RGBColor(187, 247, 208)  # #BBF7D0
    
    C_BLUE_700  = RGBColor(3, 105, 161)    # #0369A1
    C_BLUE_50   = RGBColor(239, 246, 255)  # #EFF6FF
    C_BLUE_200  = RGBColor(191, 219, 254)  # #BFDBFE

    C_RED_700   = RGBColor(185, 28, 28)    # #B91C1C
    C_RED_50    = RGBColor(254, 242, 242)  # #FEF2F2
    C_RED_200   = RGBColor(254, 202, 202)  # #FECACA

    C_WHITE     = RGBColor(255, 255, 255)
    C_YELLOW_50 = RGBColor(254, 252, 232)  # #FEFCE8

    def add_header_and_footer(slide, slide_num, total_slides=6):
        # Header background bar
        hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.55))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = C_SLATE_900
        hdr.line.color.rgb = C_SLATE_900
        
        # Amber accent line
        acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.55), Inches(13.333), Inches(0.04))
        acc.fill.solid()
        acc.fill.fore_color.rgb = C_AMBER_500
        acc.line.color.rgb = C_AMBER_500

        # Header text (Left)
        tx_hdr = slide.shapes.add_textbox(Inches(0.5), Inches(0.08), Inches(8.5), Inches(0.4))
        p = tx_hdr.text_frame.paragraphs[0]
        p.text = "SMART INDIA HACKATHON 2026  •  SOVEREIGN ON-PREMISE AGENTIC AI WORKBENCH"
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

        # Header text (Right)
        tx_hdr_r = slide.shapes.add_textbox(Inches(9.2), Inches(0.08), Inches(3.6), Inches(0.4))
        p2 = tx_hdr_r.text_frame.paragraphs[0]
        p2.text = "PS ID: SIH26117  |  MRPL"
        p2.alignment = PP_ALIGN.RIGHT
        p2.font.name = "Arial"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = C_AMBER_500

        # Footer line
        ft_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.02))
        ft_line.fill.solid()
        ft_line.fill.fore_color.rgb = C_SLATE_300
        ft_line.line.color.rgb = C_SLATE_300

        # Footer text (Left)
        tx_ft_l = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(8.5), Inches(0.35))
        pf_l = tx_ft_l.text_frame.paragraphs[0]
        pf_l.text = "@SIH Idea submission - Template  |  Mangalore Refinery and Petrochemicals Limited"
        pf_l.font.name = "Arial"
        pf_l.font.size = Pt(9.5)
        pf_l.font.color.rgb = C_AMBER_600
        pf_l.font.bold = True

        # Footer text (Right)
        tx_ft_r = slide.shapes.add_textbox(Inches(10.0), Inches(7.1), Inches(2.8), Inches(0.35))
        pf_r = tx_ft_r.text_frame.paragraphs[0]
        pf_r.text = f"Slide {slide_num} of {total_slides}"
        pf_r.alignment = PP_ALIGN.RIGHT
        pf_r.font.name = "Arial"
        pf_r.font.size = Pt(9.5)
        pf_r.font.color.rgb = C_SLATE_500
        pf_r.font.bold = True

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    add_header_and_footer(s1, 1)

    # Title & Subtitle banner
    tx_title = s1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(1.1))
    tf1 = tx_title.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "SMART INDIA HACKATHON 2026"
    p1.font.name = "Arial"
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = C_AMBER_600
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "SOVEREIGN ON-PREMISE AGENTIC AI WORKBENCH"
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = C_SLATE_900
    p2.alignment = PP_ALIGN.CENTER

    # Information Card Table
    t_shape = s1.shapes.add_table(6, 2, Inches(0.8), Inches(2.0), Inches(11.733), Inches(3.6))
    t1 = t_shape.table
    t1.columns[0].width = Inches(3.2)
    t1.columns[1].width = Inches(8.533)

    meta_items = [
        ("Problem Statement ID:", "SIH26117", True),
        ("Problem Statement Title:", "Sovereign On-Premise Agentic AI Workbench using Open-Weight Multimodal LLMs for Confidential Industrial Work", False),
        ("Organization:", "Mangalore Refinery and Petrochemicals Limited (MRPL)", True),
        ("Theme & PS Category:", "Smart Automation  |  Software", False),
        ("Team ID & Name:", "[Registered Team ID]  —  [Registered Team Name]", False),
        ("Platform & Architecture:", "Air-Gapped, Zero-Install Portable Windows/Linux Executable with Embedded Ollama Runtime", False),
    ]

    for row_idx, (label, val, highlight) in enumerate(meta_items):
        cell_lbl = t1.cell(row_idx, 0)
        cell_lbl.fill.solid()
        cell_lbl.fill.fore_color.rgb = C_SLATE_100
        p_lbl = cell_lbl.text_frame.paragraphs[0]
        p_lbl.text = label
        p_lbl.font.name = "Arial"
        p_lbl.font.size = Pt(11)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = C_SLATE_900

        cell_val = t1.cell(row_idx, 1)
        cell_val.fill.solid()
        cell_val.fill.fore_color.rgb = C_SLATE_50
        p_val = cell_val.text_frame.paragraphs[0]
        p_val.text = val
        p_val.font.name = "Arial"
        p_val.font.size = Pt(11)
        p_val.font.bold = highlight
        p_val.font.color.rgb = C_AMBER_600 if highlight else C_SLATE_800

    # Scope Summary Box at bottom
    box_s1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(1.05))
    box_s1.fill.solid()
    box_s1.fill.fore_color.rgb = C_AMBER_100
    box_s1.line.color.rgb = C_AMBER_500
    box_s1.line.width = Pt(1.5)

    tf_box = box_s1.text_frame
    tf_box.word_wrap = True
    p_box = tf_box.paragraphs[0]
    p_box.text = "Project Scope Summary:"
    p_box.font.name = "Arial"
    p_box.font.size = Pt(10.5)
    p_box.font.bold = True
    p_box.font.color.rgb = C_AMBER_600

    p_box_desc = tf_box.add_paragraph()
    p_box_desc.text = "An industrial-grade, fully sovereign on-premise desktop AI workbench built for MRPL. Delivers automated multi-turn engineering reasoning, SOP compliance audits (Hydrocracker, FCCU), P&ID drawing analysis, hardware-aware 32-model management, and 1-click executive document generation under strict zero-external-network isolation."
    p_box_desc.font.name = "Arial"
    p_box_desc.font.size = Pt(10)
    p_box_desc.font.color.rgb = C_SLATE_800

    # =========================================================================
    # SLIDE 2: Problem & Proposed Solution
    # =========================================================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header_and_footer(s2, 2)

    # Slide Title
    tx_s2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.5))
    p_s2 = tx_s2.text_frame.paragraphs[0]
    p_s2.text = "SOVEREIGN AI WORKBENCH — PROBLEM & SOLUTION"
    p_s2.font.name = "Arial"
    p_s2.font.size = Pt(18)
    p_s2.font.bold = True
    p_s2.font.color.rgb = C_SLATE_900

    # Left Column: Problem & Proposed Solution Cards
    card_prob = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.6), Inches(2.6))
    card_prob.fill.solid()
    card_prob.fill.fore_color.rgb = C_RED_50
    card_prob.line.color.rgb = C_RED_200
    card_prob.line.width = Pt(1.5)

    tf_prob = card_prob.text_frame
    tf_prob.word_wrap = True
    p_pr_h = tf_prob.paragraphs[0]
    p_pr_h.text = "THE PROBLEM"
    p_pr_h.font.name = "Arial"
    p_pr_h.font.size = Pt(13)
    p_pr_h.font.bold = True
    p_pr_h.font.color.rgb = C_RED_700

    prob_bullets = [
        "Confidential Data Risk: Sensitive approval notes, P&IDs, Hydrocracker / FCCU telemetry logs, and internal operating code cannot leave refinery premises.",
        "Cloud AI Policy Gap: Commercial cloud AI assistants (OpenAI, Claude) create catastrophic data leakage and cybersecurity compliance violations.",
        "Operational Bottleneck: Manual log verification and cross-referencing against heavy SOP manuals delays maintenance clearance and reduces plant throughput."
    ]
    for b in prob_bullets:
        pb = tf_prob.add_paragraph()
        pb.text = "• " + b
        pb.font.name = "Arial"
        pb.font.size = Pt(9.5)
        pb.font.color.rgb = C_SLATE_800

    card_sol = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.1), Inches(5.6), Inches(2.75))
    card_sol.fill.solid()
    card_sol.fill.fore_color.rgb = C_GREEN_50
    card_sol.line.color.rgb = C_GREEN_200
    card_sol.line.width = Pt(1.5)

    tf_sol = card_sol.text_frame
    tf_sol.word_wrap = True
    p_sol_h = tf_sol.paragraphs[0]
    p_sol_h.text = "THE PROPOSED SOLUTION"
    p_sol_h.font.name = "Arial"
    p_sol_h.font.size = Pt(13)
    p_sol_h.font.bold = True
    p_sol_h.font.color.rgb = C_GREEN_700

    sol_bullets = [
        "100% Self-Hosted & Sovereign: Fully offline desktop workstation running on refinery hardware with verifiable zero external network traffic.",
        "Hardware-Aware Intelligence: Auto-detects real CPU/RAM/VRAM and scores a 32-model catalog to run the best-fit model comfortably.",
        "Multi-Format File Ingestion: Native extraction from Word, Excel telemetry sheets, PDFs, code files, and P&ID diagrams.",
        "Automated Deliverable Dispatch: One-click generation of signed executive approval notes in DOCX and PDF directly to Desktop."
    ]
    for b in sol_bullets:
        pb = tf_sol.add_paragraph()
        pb.text = "• " + b
        pb.font.name = "Arial"
        pb.font.size = Pt(9.5)
        pb.font.color.rgb = C_SLATE_800

    # Right Column: 6 Core Agentic Pillars Card
    card_pillars = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.3), Inches(5.933), Inches(5.55))
    card_pillars.fill.solid()
    card_pillars.fill.fore_color.rgb = C_BLUE_50
    card_pillars.line.color.rgb = C_BLUE_200
    card_pillars.line.width = Pt(1.5)

    tf_pil = card_pillars.text_frame
    tf_pil.word_wrap = True
    p_pil_h = tf_pil.paragraphs[0]
    p_pil_h.text = "CORE AGENTIC PILLARS & WORKBENCH CAPABILITIES"
    p_pil_h.font.name = "Arial"
    p_pil_h.font.size = Pt(13)
    p_pil_h.font.bold = True
    p_pil_h.font.color.rgb = C_BLUE_700

    pillars = [
        ("1. Multi-Step Industrial Planning:", "AI audits live telemetry against MRPL SOP limits (MAWP <= 150 bar, 6-month PSV calibration window) and formats structured approval notes."),
        ("2. Local Tool & File Integration:", "Universal parser extracts text, tables, and sensor metrics from .docx, .xlsx, .pdf, code, and telemetry logs."),
        ("3. Open-Weight 32-Model Hub:", "Categorized repository (Coding, DeepSeek Reasoning, Vision, General SOP, Enterprise) with 1-click streaming downloads and cancel support."),
        ("4. Scan & Drawing Understanding:", "Multimodal visual models (LLaVA, Moondream, Llama-3.2-Vision) inspect P&ID schematics and equipment gauge readings."),
        ("5. VRAM & Process Safety Sandbox:", "Windows Job Object (KILL_ON_JOB_CLOSE), orphan cleanup, and keep_alive=0 dynamic memory release."),
        ("6. Executive Deliverable Export:", "Styled ReportLab PDF and python-docx approval note generator dispatched directly to OS Desktop.")
    ]
    for title, desc in pillars:
        pt = tf_pil.add_paragraph()
        pt.text = title
        pt.font.name = "Arial"
        pt.font.size = Pt(10)
        pt.font.bold = True
        pt.font.color.rgb = C_SLATE_900

        pd = tf_pil.add_paragraph()
        pd.text = desc
        pd.font.name = "Arial"
        pd.font.size = Pt(9)
        pd.font.color.rgb = C_SLATE_700

    # =========================================================================
    # SLIDE 3: Technical Approach
    # =========================================================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header_and_footer(s3, 3)

    tx_s3 = s3.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.5))
    p_s3 = tx_s3.text_frame.paragraphs[0]
    p_s3.text = "TECHNICAL APPROACH & ARCHITECTURE"
    p_s3.font.name = "Arial"
    p_s3.font.size = Pt(18)
    p_s3.font.bold = True
    p_s3.font.color.rgb = C_SLATE_900

    # Left: Layered Stack Card
    card_stack = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(4.5))
    card_stack.fill.solid()
    card_stack.fill.fore_color.rgb = C_SLATE_50
    card_stack.line.color.rgb = C_SLATE_300
    card_stack.line.width = Pt(1.5)

    tf_st = card_stack.text_frame
    tf_st.word_wrap = True
    p_st_h = tf_st.paragraphs[0]
    p_st_h.text = "1. LAYERED LOCAL WORKBENCH STACK"
    p_st_h.font.name = "Arial"
    p_st_h.font.size = Pt(12)
    p_st_h.font.bold = True
    p_st_h.font.color.rgb = C_SLATE_900

    stack_layers = [
        ("Presentation Layer:", "Single-Page App in PyWebView (HTML5/CSS3/JS); 5 luxury themes (Midnight Obsidian, Titanium, Azure, Matrix, Amethyst) + custom palette builder; real-time stop button."),
        ("Backend Server (main.py):", "Multithreaded TCP Server on port 8085; REST API + NDJSON streaming progress endpoints; daemon thread dispatch."),
        ("Core Engine (sovereign_engine.py):", "Hardware detection, 32-model catalog scorer, multi-turn chat orchestrator, universal file parser."),
        ("Process & Memory Supervisor:", "Win32 Job Object (0x2000), startup orphan process killer, dynamic VRAM model unloader (keep_alive=0)."),
        ("Inference Layer:", "Embedded portable Ollama daemon (port 11434) running quantized GGUF weights on CUDA GPU / CPU AVX2."),
        ("Air-Gap Isolation:", "Hardcoded loopback (127.0.0.1), zero outbound calls, standalone data/ directory structure.")
    ]
    for l_title, l_desc in stack_layers:
        pt = tf_st.add_paragraph()
        pt.text = "• " + l_title + " " + l_desc
        pt.font.name = "Arial"
        pt.font.size = Pt(8.8)
        pt.font.color.rgb = C_SLATE_800

    # Right: Inspection Workflow Card
    card_wf = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.3), Inches(5.833), Inches(4.5))
    card_wf.fill.solid()
    card_wf.fill.fore_color.rgb = C_BLUE_50
    card_wf.line.color.rgb = C_BLUE_200
    card_wf.line.width = Pt(1.5)

    tf_wf = card_wf.text_frame
    tf_wf.word_wrap = True
    p_wf_h = tf_wf.paragraphs[0]
    p_wf_h.text = "2. STREAMLINED INSPECTION & APPROVAL WORKFLOW"
    p_wf_h.font.name = "Arial"
    p_wf_h.font.size = Pt(12)
    p_wf_h.font.bold = True
    p_wf_h.font.color.rgb = C_BLUE_700

    wf_steps = [
        ("Step 1: Telemetry & Report Ingestion:", "Operator attaches Hydrocracker / FCCU logs, Excel sensor data, or P&ID drawings."),
        ("Step 2: Context Parsing:", "Universal parser extracts operating pressures (142.5 bar), temperatures (410°C), and valve tags (PSV-102A)."),
        ("Step 3: Hardware-Guided Model Routing:", "Engine matches prompt complexity to the optimal installed model (e.g., DeepSeek-R1 / Qwen2.5)."),
        ("Step 4: SOP Compliance Audit:", "LLM checks operating pressure against MRPL SOP limits (MAWP: 150 bar, 140 bar threshold audit)."),
        ("Step 5: Executive Deliverable Generation:", "Auto-formats structured findings with executive summary, compliance status, and signature lines."),
        ("Step 6: Direct Desktop Dispatch:", "Builds formatted .docx and styled .pdf documents directly on Desktop and launches in Word.")
    ]
    for s_title, s_desc in wf_steps:
        pt = tf_wf.add_paragraph()
        pt.text = s_title
        pt.font.name = "Arial"
        pt.font.size = Pt(9.5)
        pt.font.bold = True
        pt.font.color.rgb = C_SLATE_900

        pd = tf_wf.add_paragraph()
        pd.text = s_desc
        pd.font.name = "Arial"
        pd.font.size = Pt(8.8)
        pd.font.color.rgb = C_SLATE_700

    # Bottom Implementation Banner
    banner_s3 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.9))
    banner_s3.fill.solid()
    banner_s3.fill.fore_color.rgb = C_SLATE_100
    banner_s3.line.color.rgb = C_SLATE_300

    tf_b3 = banner_s3.text_frame
    tf_b3.word_wrap = True
    p_b3 = tf_b3.paragraphs[0]
    p_b3.text = "Implementation Components:"
    p_b3.font.name = "Arial"
    p_b3.font.size = Pt(9.5)
    p_b3.font.bold = True
    p_b3.font.color.rgb = C_AMBER_600

    p_b3_desc = tf_b3.add_paragraph()
    p_b3_desc.text = "Agent Harness: Multithreaded Python (main.py + sovereign_engine.py) | Inference: Embedded Portable Ollama + GGUF Quantized Models (CUDA / CPU AVX2) | File Ingestion: Native python-docx, openpyxl, PDF & text parsers | UI Framework: PyWebView + Luxury Responsive HTML5 SPA (5 Themes) | Safety & Isolation: Win32 Job Object (KILL_ON_JOB_CLOSE) + Dynamic VRAM Unloader (keep_alive: 0) + Zero-Egress Loopback"
    p_b3_desc.font.name = "Arial"
    p_b3_desc.font.size = Pt(8.5)
    p_b3_desc.font.color.rgb = C_SLATE_800

    # =========================================================================
    # SLIDE 4: Feasibility and Viability
    # =========================================================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_header_and_footer(s4, 4)

    tx_s4 = s4.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.5))
    p_s4 = tx_s4.text_frame.paragraphs[0]
    p_s4.text = "FEASIBILITY, VIABILITY & RISK MITIGATION"
    p_s4.font.name = "Arial"
    p_s4.font.size = Pt(18)
    p_s4.font.bold = True
    p_s4.font.color.rgb = C_SLATE_900

    # Stepper Horizontal Flow
    steps = [
        ("1. Foundation", "Portable EXE, server & embedded engine"),
        ("2. Model Hub", "32 models, 1-click streaming pull"),
        ("3. Guardrails", "Job Object, VRAM unloader & orphan killer"),
        ("4. Agentic Core", "Multi-turn memory & MRPL SOP prompt"),
        ("5. File Ingestion", "DOCX, XLSX, PDF & P&ID parser"),
        ("6. Deliverables", "ReportLab PDF & Word to Desktop")
    ]
    step_w = Inches(1.85)
    for idx, (st_t, st_d) in enumerate(steps):
        bx = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 1.98), Inches(1.3), step_w, Inches(0.95))
        bx.fill.solid()
        bx.fill.fore_color.rgb = C_SLATE_100
        bx.line.color.rgb = C_SLATE_300
        
        tfx = bx.text_frame
        tfx.word_wrap = True
        px1 = tfx.paragraphs[0]
        px1.text = st_t
        px1.font.name = "Arial"
        px1.font.size = Pt(9.5)
        px1.font.bold = True
        px1.font.color.rgb = C_AMBER_600
        px1.alignment = PP_ALIGN.CENTER

        px2 = tfx.add_paragraph()
        px2.text = st_d
        px2.font.name = "Arial"
        px2.font.size = Pt(7.5)
        px2.font.color.rgb = C_SLATE_700
        px2.alignment = PP_ALIGN.CENTER

    # Left: Feasibility Table
    t_feas_shape = s4.shapes.add_table(5, 2, Inches(0.8), Inches(2.45), Inches(5.6), Inches(4.35))
    t_feas = t_feas_shape.table
    t_feas.columns[0].width = Inches(1.8)
    t_feas.columns[1].width = Inches(3.8)

    feas_data = [
        ("Single Workstation Feasibility", "Implementation & Verification"),
        ("Open-Weight Models", "Curated 32-model catalog (0.135B to 70B) covering Qwen2.5, DeepSeek-R1, Llama-3.2, Gemma-2, Moondream, Mistral."),
        ("Quantized GGUF", "4-bit/8-bit quantization fits low-end GPUs (2GB VRAM GT 710) up to multi-GPU high-memory servers."),
        ("Embedded Runtime", "Standalone Ollama daemon with isolated storage (data/models/); zero admin rights or installs needed."),
        ("Portable Storage", "All settings, sessions, models, and exports stored in portable data/ folder packaged in single ~21MB ZIP.")
    ]
    for r_idx, (c1, c2) in enumerate(feas_data):
        cell1, cell2 = t_feas.cell(r_idx, 0), t_feas.cell(r_idx, 1)
        if r_idx == 0:
            cell1.fill.solid(); cell1.fill.fore_color.rgb = C_SLATE_900
            cell2.fill.solid(); cell2.fill.fore_color.rgb = C_SLATE_900
            p_c1 = cell1.text_frame.paragraphs[0]; p_c1.text = c1; p_c1.font.bold = True; p_c1.font.color.rgb = C_WHITE; p_c1.font.size = Pt(9.5)
            p_c2 = cell2.text_frame.paragraphs[0]; p_c2.text = c2; p_c2.font.bold = True; p_c2.font.color.rgb = C_WHITE; p_c2.font.size = Pt(9.5)
        else:
            cell1.fill.solid(); cell1.fill.fore_color.rgb = C_SLATE_100
            cell2.fill.solid(); cell2.fill.fore_color.rgb = C_SLATE_50
            p_c1 = cell1.text_frame.paragraphs[0]; p_c1.text = c1; p_c1.font.bold = True; p_c1.font.color.rgb = C_SLATE_900; p_c1.font.size = Pt(9)
            p_c2 = cell2.text_frame.paragraphs[0]; p_c2.text = c2; p_c2.font.color.rgb = C_SLATE_800; p_c2.font.size = Pt(8.5)

    # Right: Risk & Mitigation Table
    t_risk_shape = s4.shapes.add_table(5, 3, Inches(6.6), Inches(2.45), Inches(5.933), Inches(4.35))
    t_risk = t_risk_shape.table
    t_risk.columns[0].width = Inches(1.7)
    t_risk.columns[1].width = Inches(2.1)
    t_risk.columns[2].width = Inches(2.133)

    risk_data = [
        ("Identified Risk", "Control Implemented", "Live Demo Evidence"),
        ("Model Exceeds VRAM (OOM)", "Hardware profiler + formula: Mem = Params * 0.65 + 0.8 GB", "Dynamic badges (Optimal, Full GPU, CPU) prevent OOM crashes."),
        ("VRAM Leak / Zombie Tasks", "Win32 Job Object (0x2000) + dynamic keep_alive: 0 unloader", "Verified via nvidia-smi: 0 MB VRAM retained after exit/switch."),
        ("Cloud Data Leakage", "Hardcoded loopback (127.0.0.1) with zero external calls", "Wireshark records 0 outbound packets during chat and export."),
        ("Interrupted Downloads", "NDJSON streaming reader with thread cancellation event", "Live byte progress bar with instant responsive Cancel button.")
    ]
    for r_idx, (c1, c2, c3) in enumerate(risk_data):
        cell1, cell2, cell3 = t_risk.cell(r_idx, 0), t_risk.cell(r_idx, 1), t_risk.cell(r_idx, 2)
        if r_idx == 0:
            for cell, text in [(cell1, c1), (cell2, c2), (cell3, c3)]:
                cell.fill.solid(); cell.fill.fore_color.rgb = C_AMBER_600
                p_c = cell.text_frame.paragraphs[0]; p_c.text = text; p_c.font.bold = True; p_c.font.color.rgb = C_WHITE; p_c.font.size = Pt(9.5)
        else:
            for cell, text, is_bold in [(cell1, c1, True), (cell2, c2, False), (cell3, c3, False)]:
                cell.fill.solid(); cell.fill.fore_color.rgb = C_SLATE_50
                p_c = cell.text_frame.paragraphs[0]; p_c.text = text; p_c.font.bold = is_bold; p_c.font.color.rgb = C_SLATE_900 if is_bold else C_SLATE_800; p_c.font.size = Pt(8.5)

    # =========================================================================
    # SLIDE 5: Impact and Benefits
    # =========================================================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_header_and_footer(s5, 5)

    tx_s5 = s5.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.5))
    p_s5 = tx_s5.text_frame.paragraphs[0]
    p_s5.text = "IMPACT, BENEFITS & INDUSTRIAL VALUE"
    p_s5.font.name = "Arial"
    p_s5.font.size = Pt(18)
    p_s5.font.bold = True
    p_s5.font.color.rgb = C_SLATE_900

    # Left: 4-Gear Process & Why It Matters
    card_wf_imp = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.6), Inches(5.55))
    card_wf_imp.fill.solid()
    card_wf_imp.fill.fore_color.rgb = C_SLATE_50
    card_wf_imp.line.color.rgb = C_SLATE_300
    card_wf_imp.line.width = Pt(1.5)

    tf_imp = card_wf_imp.text_frame
    tf_imp.word_wrap = True
    p_imp_h = tf_imp.paragraphs[0]
    p_imp_h.text = "WORKFLOW TRANSFORMATION"
    p_imp_h.font.name = "Arial"
    p_imp_h.font.size = Pt(13)
    p_imp_h.font.bold = True
    p_imp_h.font.color.rgb = C_SLATE_900

    gears = [
        ("⚙️ Gear 1: Profile & Route:", "Hardware profiler auto-detects GPU/RAM and routes task to optimal model from 32-model hub."),
        ("⚙️ Gear 2: Read & Ingest:", "Universal parser extracts telemetry metrics from Excel sensor logs, DOCX, and P&ID drawings."),
        ("⚙️ Gear 3: Audit in Sandbox:", "Local LLM evaluates operational parameters against MRPL SOP limits (MAWP threshold, valve dates)."),
        ("⚙️ Gear 4: Deliver & Clear:", "Exports signed executive approval note to Desktop and releases VRAM immediately.")
    ]
    for g_t, g_d in gears:
        pt = tf_imp.add_paragraph()
        pt.text = g_t
        pt.font.name = "Arial"
        pt.font.size = Pt(9.5)
        pt.font.bold = True
        pt.font.color.rgb = C_SLATE_900

        pd = tf_imp.add_paragraph()
        pd.text = g_d
        pd.font.name = "Arial"
        pd.font.size = Pt(9)
        pd.font.color.rgb = C_SLATE_700

    p_w_h = tf_imp.add_paragraph()
    p_w_h.text = "WHY IT MATTERS"
    p_w_h.font.name = "Arial"
    p_w_h.font.size = Pt(12)
    p_w_h.font.bold = True
    p_w_h.font.color.rgb = C_AMBER_600

    whys = [
        "Absolute Confidentiality: Telemetry logs, tariff accounting, and internal memos never leave refinery boundaries.",
        "Zero-Friction Portability: Single ~21MB ZIP runs out-of-the-box on Windows/Linux without admin privileges or pip installs.",
        "Hardware Inclusivity: Runs on modest 2GB VRAM cards (GT 710) up to multi-GPU enterprise AI workstations."
    ]
    for w in whys:
        pw = tf_imp.add_paragraph()
        pw.text = "• " + w
        pw.font.name = "Arial"
        pw.font.size = Pt(8.8)
        pw.font.color.rgb = C_SLATE_800

    # Right: Stakeholder Matrix Card
    card_stk = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.3), Inches(5.933), Inches(5.55))
    card_stk.fill.solid()
    card_stk.fill.fore_color.rgb = C_YELLOW_50
    card_stk.line.color.rgb = C_AMBER_500
    card_stk.line.width = Pt(1.5)

    tf_stk = card_stk.text_frame
    tf_stk.word_wrap = True
    p_stk_h = tf_stk.paragraphs[0]
    p_stk_h.text = "STAKEHOLDER VALUE MATRIX (BEFORE VS. AFTER)"
    p_stk_h.font.name = "Arial"
    p_stk_h.font.size = Pt(13)
    p_stk_h.font.bold = True
    p_stk_h.font.color.rgb = C_AMBER_600

    stk_groups = [
        ("👨‍💼 Plant & Operations Engineers:",
         "Before: Manual spreadsheet parsing & cross-referencing against 200+ page SOP manuals.",
         "After: Instant file ingestion, automated telemetry audits, and evidence-backed draft notes."),
        ("🛡️ Operations & Safety Approvers:",
         "Before: Slow, multi-day document review cycles causing operational delays.",
         "After: Structured, standardized, and traceable compliance memos generated in seconds."),
        ("🔒 IT & Cybersecurity Officers:",
         "Before: Massive risk of proprietary refinery data leakage via public cloud LLMs.",
         "After: 100% on-premise air-gapped execution with verifiable zero egress traffic."),
        ("🏢 Executive Management:",
         "Before: High recurring cloud API costs, vendor lock-in, and compliance penalties.",
         "After: Zero license costs, offline USB-deployable solution across all refinery units.")
    ]
    for grp_t, grp_b, grp_a in stk_groups:
        pt = tf_stk.add_paragraph()
        pt.text = grp_t
        pt.font.name = "Arial"
        pt.font.size = Pt(9.5)
        pt.font.bold = True
        pt.font.color.rgb = C_SLATE_900

        pb = tf_stk.add_paragraph()
        pb.text = grp_b
        pb.font.name = "Arial"
        pb.font.size = Pt(8.5)
        pb.font.color.rgb = C_SLATE_700

        pa = tf_stk.add_paragraph()
        pa.text = grp_a
        pa.font.name = "Arial"
        pa.font.size = Pt(8.5)
        pa.font.bold = True
        pa.font.color.rgb = C_GREEN_700

    # =========================================================================
    # SLIDE 6: Research, References & Implementation Roadmap
    # =========================================================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_header_and_footer(s6, 6)

    tx_s6 = s6.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.5))
    p_s6 = tx_s6.text_frame.paragraphs[0]
    p_s6.text = "RESEARCH, REFERENCES & IMPLEMENTATION ROADMAP"
    p_s6.font.name = "Arial"
    p_s6.font.size = Pt(18)
    p_s6.font.bold = True
    p_s6.font.color.rgb = C_SLATE_900

    # Left: Demonstration Checklist Card
    card_chk = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.6), Inches(4.5))
    card_chk.fill.solid()
    card_chk.fill.fore_color.rgb = C_GREEN_50
    card_chk.line.color.rgb = C_GREEN_200
    card_chk.line.width = Pt(1.5)

    tf_chk = card_chk.text_frame
    tf_chk.word_wrap = True
    p_chk_h = tf_chk.paragraphs[0]
    p_chk_h.text = "LIVE DEMONSTRATION CHECKLIST"
    p_chk_h.font.name = "Arial"
    p_chk_h.font.size = Pt(12)
    p_chk_h.font.bold = True
    p_chk_h.font.color.rgb = C_GREEN_700

    chks = [
        ("[x] Hardware Auto-Profiling:", "Accurately detects CPU cores, RAM, free disk, and GPU VRAM via nvidia-smi and auto-recommends optimal model."),
        ("[x] Interactive 32-Model Hub:", "1-click NDJSON streaming download with real-time progress bar and responsive cancel support."),
        ("[x] Multi-Format Context Ingestion:", "Extracts and injects text from Word documents, Excel telemetry sheets, PDFs, and code files into prompt."),
        ("[x] Executive Deliverable Dispatch:", "One-click generation of styled ReportLab PDFs and formatted Microsoft Word (.docx) approval memos to Desktop."),
        ("[x] 100% VRAM & Process Safety:", "Windows Job Object (KILL_ON_JOB_CLOSE) + keep_alive: 0 model unloading ensures 0 MB VRAM leakage.")
    ]
    for c_t, c_d in chks:
        pt = tf_chk.add_paragraph()
        pt.text = c_t + " " + c_d
        pt.font.name = "Arial"
        pt.font.size = Pt(9)
        pt.font.color.rgb = C_SLATE_800

    # Right: 10-Stage Roadmap Card
    card_rdm = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.3), Inches(5.933), Inches(4.5))
    card_rdm.fill.solid()
    card_rdm.fill.fore_color.rgb = C_SLATE_50
    card_rdm.line.color.rgb = C_SLATE_300
    card_rdm.line.width = Pt(1.5)

    tf_rdm = card_rdm.text_frame
    tf_rdm.word_wrap = True
    p_rdm_h = tf_rdm.paragraphs[0]
    p_rdm_h.text = "10-STAGE IMPLEMENTATION ROADMAP"
    p_rdm_h.font.name = "Arial"
    p_rdm_h.font.size = Pt(12)
    p_rdm_h.font.bold = True
    p_rdm_h.font.color.rgb = C_SLATE_900

    stages = [
        "1. SIH Scope: Defined air-gapped, multimodal sovereign AI scope for MRPL refinery.",
        "2. Portable Architecture: Built self-contained data/ folder structure and standalone EXE.",
        "3. Embedded Inference: Bundled standalone Ollama daemon with isolated OLLAMA_MODELS.",
        "4. Universal Ingestion: Native python-docx, openpyxl, PDF and diagram text extractors.",
        "5. Strict Air-Gap Proof: Hardcoded loopback (127.0.0.1) with zero outbound network calls.",
        "6. Hardware Auto-Scoring: Memory formula scoring with dynamic visual fit badges.",
        "7. Curated 32-Model Hub: Categorized catalog with 1-click download and live progress bars.",
        "8. VRAM & Process Guard: Win32 Job Object supervisor and dynamic keep_alive: 0 unloader.",
        "9. Document Generator: ReportLab PDF formatting and direct desktop Word dispatch.",
        "10. Luxury Modern UI: 5 preset themes, custom palette builder, and stop response controller."
    ]
    for stg in stages:
        pt = tf_rdm.add_paragraph()
        pt.text = stg
        pt.font.name = "Arial"
        pt.font.size = Pt(8.2)
        pt.font.color.rgb = C_SLATE_800

    # Bottom Reference Banner
    banner_s6 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.9))
    banner_s6.fill.solid()
    banner_s6.fill.fore_color.rgb = C_SLATE_100
    banner_s6.line.color.rgb = C_SLATE_300

    tf_b6 = banner_s6.text_frame
    tf_b6.word_wrap = True
    p_b6 = tf_b6.paragraphs[0]
    p_b6.text = "Technical References & URLs:"
    p_b6.font.name = "Arial"
    p_b6.font.size = Pt(9)
    p_b6.font.bold = True
    p_b6.font.color.rgb = C_SLATE_700

    p_b6_desc = tf_b6.add_paragraph()
    p_b6_desc.text = "SIH Portal: sih.gov.in/sih2026PS (Problem Statement SIH26117) | Inference Engine: ollama.com, github.com/ollama/ollama, github.com/ggml-org/llama.cpp | Parsing & Docs: github.com/python-openxml/python-docx, github.com/openpyxl/openpyxl, github.com/reportlab/reportlab | Open-Weight Models: github.com/qwenlm/Qwen2.5, github.com/deepseek-ai/DeepSeek-R1, github.com/meta-llama/llama-models, github.com/vikhyat/moondream"
    p_b6_desc.font.name = "Arial"
    p_b6_desc.font.size = Pt(7.8)
    p_b6_desc.font.color.rgb = C_SLATE_600

    # Save presentation
    prs.save(output_path)
    print(f"PowerPoint Presentation successfully generated at: {output_path}")

if __name__ == '__main__':
    root_dir = os.path.dirname(os.path.abspath(__file__))
    out_pptx = os.path.join(root_dir, "SIH2026_SIH26117_Sovereign_AI_Workbench_Presentation.pptx")
    create_presentation(out_pptx)
